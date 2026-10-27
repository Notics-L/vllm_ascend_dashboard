"""
报告生成器 - PDF 和 Excel 报告生成
"""
import json
import logging
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

logger = logging.getLogger(__name__)


def fmt_duration(sec: float) -> str:
    if sec < 60:
        return f"{sec:.0f}s"
    m = int(sec // 60)
    s = int(sec % 60)
    if m < 60:
        return f"{m}m{s}s" if s > 0 else f"{m}m"
    h = m // 60
    m = m % 60
    return f"{h}h{m}m"


class ReportGenerator:
    """报告生成器"""

    # 样式常量
    title_font = Font(name="Microsoft YaHei", bold=True, size=16, color="1F4E79")
    section_font = Font(name="Microsoft YaHei", bold=True, size=12, color="2E75B6")
    header_font = Font(name="Microsoft YaHei", bold=True, size=10, color="FFFFFF")
    normal_font = Font(name="Microsoft YaHei", size=9)
    bold_font = Font(name="Microsoft YaHei", bold=True, size=9)
    red_bold = Font(name="Microsoft YaHei", bold=True, size=10, color="C00000")

    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    success_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    failure_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    warn_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    @staticmethod
    def _style_header(ws, row: int, ncols: int):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = ReportGenerator.header_font
            cell.fill = ReportGenerator.header_fill
            cell.alignment = ReportGenerator.header_align
            cell.border = ReportGenerator.thin_border

    @staticmethod
    def _status_fill(status: str):
        s = str(status).lower()
        if s == "success":
            return ReportGenerator.success_fill
        elif s == "failure":
            return ReportGenerator.failure_fill
        elif s in ("skipped", "cancelled", "action_required"):
            return ReportGenerator.warn_fill
        return None

    @staticmethod
    def _style_row(ws, row: int, ncols: int):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = ReportGenerator.normal_font
            cell.border = ReportGenerator.thin_border
            cell.alignment = ReportGenerator.center_align

    async def generate_excel(
        self,
        workflow_name: str,
        stats: dict[str, Any],
        runs: list[dict[str, Any]],
        groups: list[dict[str, Any]],
        evaluation: dict[str, Any] | None = None,
    ) -> bytes:
        wb = Workbook()

        # Sheet 1: 综合概览
        self._build_overview_sheet(wb, workflow_name, stats, evaluation)

        # Sheet 2: 构建记录明细
        self._build_runs_sheet(wb, runs)

        # Sheet 3: Job Group 层级
        self._build_groups_sheet(wb, groups)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    def _build_overview_sheet(self, wb: Workbook, workflow_name: str,
                               stats: dict, evaluation: dict | None):
        ws = wb.active
        ws.title = "综合概览"
        ws.sheet_properties.tabColor = "2E75B6"

        ws.merge_cells("A1:F1")
        ws["A1"] = f"{workflow_name} 构建分析报告"
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 35

        ws["A3"] = "基本统计"
        ws["A3"].font = self.section_font
        ws.merge_cells("A3:F3")

        stats_list = [
            ("生成时间", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S")),
            ("运行总数", str(stats.get("total", 0))),
            ("成功次数", str(stats.get("successful", 0))),
            ("失败次数", str(stats.get("failed", 0))),
            ("成功率", f"{stats.get('success_rate', 0):.1f}%"),
            ("平均执行耗时", fmt_duration(stats.get("avg_exec", 0))),
            ("最短执行耗时", fmt_duration(stats.get("min_exec", 0))),
            ("最长执行耗时", fmt_duration(stats.get("max_exec", 0))),
            ("耗时波动系数 (CV)", f"{stats.get('cv_exec', 0):.1f}%"),
            ("平均排队时间", fmt_duration(stats.get("avg_wait", 0))),
            ("平均墙钟总耗时", fmt_duration(stats.get("avg_wall", 0))),
            ("异常慢构建", str(stats.get("abnormal_count", 0))),
            ("严重排队次数", str(stats.get("long_wait_count", 0))),
        ]

        for i, (label, value) in enumerate(stats_list, 4):
            cell_l = ws.cell(row=i, column=1, value=label)
            cell_l.font = self.bold_font
            cell_l.border = self.thin_border
            cell_l.fill = self.gray_fill
            cell_v = ws.cell(row=i, column=2, value=value)
            cell_v.font = self.red_bold
            cell_v.border = self.thin_border

        if evaluation:
            eval_row = len(stats_list) + 6
            ws.cell(row=eval_row, column=1, value="AI 评估摘要").font = self.section_font
            ws.merge_cells(f"A{eval_row}:F{eval_row}")

            ai_stats = [
                ("可信度评分", f"{evaluation.get('trust_score', 0):.1f}/100"),
                ("可信度等级", evaluation.get("trust_level", "N/A")),
                ("是否生产级", "是" if evaluation.get("production_ready") else "否"),
                ("执行摘要", evaluation.get("executive_summary", "N/A")),
            ]
            for i, (label, value) in enumerate(ai_stats, eval_row + 1):
                ws.cell(row=i, column=1, value=label).font = self.bold_font
                ws.cell(row=i, column=1).border = self.thin_border
                ws.cell(row=i, column=2, value=value).font = self.normal_font
                ws.cell(row=i, column=2).border = self.thin_border

    def _build_runs_sheet(self, wb: Workbook, runs: list[dict]):
        ws = wb.create_sheet("构建记录明细")

        headers = ["运行ID", "状态", "分支", "事件", "创建时间", "排队时间", "执行时间", "墙钟时间", "Job数"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=h)
        self._style_header(ws, 1, len(headers))

        for idx, r in enumerate(runs, 2):
            ws.cell(row=idx, column=1, value=r.get("github_run_id", ""))
            ws.cell(row=idx, column=2, value=r.get("status", ""))
            ws.cell(row=idx, column=3, value=r.get("branch", ""))
            ws.cell(row=idx, column=4, value=r.get("event", ""))
            ws.cell(row=idx, column=5, value=r.get("created_at", ""))
            ws.cell(row=idx, column=6, value=fmt_duration(r.get("wait_duration", 0)))
            ws.cell(row=idx, column=7, value=fmt_duration(r.get("exec_duration", 0)))
            ws.cell(row=idx, column=8, value=fmt_duration(r.get("wall_duration", 0)))
            ws.cell(row=idx, column=9, value=r.get("num_jobs", 0))

            status_fill = self._status_fill(r.get("status", ""))
            if status_fill:
                ws.cell(row=idx, column=2).fill = status_fill

            self._style_row(ws, idx, len(headers))

    def _build_groups_sheet(self, wb: Workbook, groups: list[dict]):
        ws = wb.create_sheet("Job 层级数据")

        headers = ["Group", "Target", "OS", "Phase", "状态", "总耗时(秒)", "格式化耗时", "架构", "Sub-Job", "Runner", "Steps数"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=h)
        self._style_header(ws, 1, len(headers))

        for idx, g in enumerate(groups, 2):
            ws.cell(row=idx, column=1, value=g.get("group_name", ""))
            ws.cell(row=idx, column=1).alignment = self.left_align
            ws.cell(row=idx, column=2, value=g.get("target", ""))
            ws.cell(row=idx, column=3, value=g.get("os_name", ""))
            ws.cell(row=idx, column=4, value=g.get("phase", ""))
            ws.cell(row=idx, column=5, value=g.get("status", ""))
            ws.cell(row=idx, column=6, value=g.get("total_duration", 0))
            ws.cell(row=idx, column=7, value=fmt_duration(g.get("total_duration", 0)))
            ws.cell(row=idx, column=8, value=g.get("architecture", ""))
            ws.cell(row=idx, column=9, value=g.get("sub_job_name", ""))
            ws.cell(row=idx, column=9).alignment = self.left_align
            ws.cell(row=idx, column=10, value=g.get("runner_name", ""))
            ws.cell(row=idx, column=11, value=len(g.get("steps", [])))

            status_fill = self._status_fill(g.get("status", ""))
            if status_fill:
                ws.cell(row=idx, column=5).fill = status_fill

            self._style_row(ws, idx, len(headers))

    async def generate_pdf(
        self,
        workflow_name: str,
        stats: dict[str, Any],
        runs: list[dict[str, Any]],
        groups: list[dict[str, Any]],
        evaluation: dict[str, Any] | None = None,
    ) -> tuple[bytes, str]:
        """生成 PDF 分析报告"""
        from jinja2 import Template
        from weasyprint import HTML

        import importlib.resources as resources

        template_path = resources.files("release_management").joinpath("pdf_template.html")
        template_source = template_path.read_text(encoding="utf-8")
        template = Template(template_source)

        def _status_label(status: str) -> str:
            mapping = {
                "success": "成功",
                "failure": "失败",
                "cancelled": "取消",
                "skipped": "跳过",
                "running": "运行中",
                "action_required": "需处理",
            }
            return mapping.get(str(status).lower(), status)

        html_content = template.render(
            workflow_name=workflow_name,
            generated_at=datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S"),
            stats=stats,
            runs=runs[:15],
            groups=groups[:30],
            evaluation=evaluation,
            fmt_duration=fmt_duration,
            status_label=_status_label,
        )

        pdf_bytes = HTML(string=html_content).write_pdf()
        ts = datetime.now(UTC).strftime("%Y%m%d%H%M%S")
        filename = f"release_build_report_{ts}.pdf"
        return pdf_bytes, filename
