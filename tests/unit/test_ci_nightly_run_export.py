from io import BytesIO

from openpyxl import load_workbook

from datetime import datetime

from api.v1.ci import _build_nightly_runs_xlsx, _nightly_export_run_time_bounds


def test_nightly_run_export_dates_use_shanghai_day_boundaries():
    start, end = _nightly_export_run_time_bounds("2026-10-15", "2026-10-15")

    assert start == datetime(2026, 10, 14, 16, 0, 0)
    assert end == datetime(2026, 10, 15, 16, 0, 0)


def test_nightly_runs_xlsx_is_openpyxl_readable_with_links_and_merges():
    content = _build_nightly_runs_xlsx([
        {
            "time": "2026-10-15 09:00:00",
            "test_content": "single-node (main, Model-A)",
            "model": "Model-A",
            "coverage": "A3 / 单机",
            "log_url": "https://github.com/vllm-project/vllm-ascend/actions/runs/1/job/2",
            "model_fo": "FO-A",
            "ai_log_url": "https://dashboard.example/api/v1/ci/public/analysis/token/report",
        },
        {
            "time": "2026-10-15 09:00:00",
            "test_content": "double-node (main, Model-B)",
            "model": "Model-B",
            "coverage": "A3 / 双机",
            "log_url": "",
            "model_fo": "FO-B",
            "ai_log_url": "",
        },
    ])

    workbook = load_workbook(BytesIO(content))
    worksheet = workbook["用例运行明细"]

    assert worksheet.title == "用例运行明细"
    assert [cell.value for cell in worksheet[1]] == [
        "时间", "测试内容", "模型", "覆盖机型", "日志链接", "模型 FO", "AI 日志链接",
    ]
    assert worksheet.max_row == 3
    assert worksheet["B2"].value == "single-node (main, Model-A)"
    assert worksheet["E2"].hyperlink is not None
    assert worksheet["E2"].hyperlink.target.endswith("/runs/1/job/2")
    assert worksheet["G2"].value == "https://dashboard.example/api/v1/ci/public/analysis/token/report"
    assert worksheet["G2"].hyperlink is not None
    assert worksheet["G2"].hyperlink.target.endswith("/analysis/token/report")
    assert "A2:A3" in {str(item) for item in worksheet.merged_cells.ranges}
